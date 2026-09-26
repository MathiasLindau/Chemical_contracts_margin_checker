"""Pull the seven free series and append one row to api_price.csv.

No header. One row per calendar day. A second run on the same day
replaces that row. The file is left unchanged when any series fails.

Column order:

    pulled_on, eurusd, eurusd_as_of, ecb_deposit, ecb_deposit_as_of,
    euribor_3m, euribor_3m_as_of, sofr, sofr_as_of, brent, brent_as_of,
    gas_eu, gas_eu_as_of, maize, maize_as_of

    python -m src.margin_checker.market_probe
"""

from __future__ import annotations

import csv
import io
import json
import re
import sys
from datetime import date
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

TIMEOUT = 30
ECB = (
    ("fx", "EURUSD", "EXR/D.USD.EUR.SP00.A", "USD/EUR", "business day"),
    ("rate", "ECB_DEPOSIT", "FM/D.U2.EUR.4F.KR.DFR.LEV", "percent", "when changed"),
    ("rate", "EURIBOR_3M", "FM/M.U2.EUR.RT.MM.EURIBOR3MD_.HSTA", "percent", "monthly"),
)
WB = (
    ("raw", "BRENT", "Crude oil, Brent", "USD/bbl"),
    ("energy", "GAS_EU", "Natural gas, Europe", "USD/mmbtu"),
    ("raw", "MAIZE", "Maize", "USD/t"),
)


def output_path():
    here = Path(__file__).resolve()
    if here.parent.name == "market":
        return here.with_name("api_price.csv")
    return here.parents[2] / "data" / "market" / "api_price.csv"


def http_get(url, headers=None):
    request = Request(url, headers={"User-Agent": "margin-checker", **(headers or {})})
    try:
        with urlopen(request, timeout=TIMEOUT) as response:
            return response.status, response.read()
    except HTTPError as exc:
        return exc.code, exc.read()
    except URLError as exc:
        return 0, str(exc.reason).encode()


def show(group, name, when, value, unit, cadence):
    print(f"{group:<8} {name:<14} {when:<12} {value:<12} {unit:<12} {cadence}", flush=True)


def ecb_last(key):
    status, body = http_get(
        "https://data-api.ecb.europa.eu/service/data/" + key + "?lastNObservations=1",
        {"Accept": "text/csv"},
    )
    if status != 200:
        raise RuntimeError("HTTP " + str(status))
    row = list(csv.DictReader(io.StringIO(body.decode())))[-1]
    return row["TIME_PERIOD"], row["OBS_VALUE"]


def fetch_ecb():
    found = {}
    for group, name, key, unit, cadence in ECB:
        when, value = ecb_last(key)
        found[name] = (when, value)
        show(group, name, when, value, unit, cadence)
    return found


def fetch_sofr():
    status, body = http_get("https://markets.newyorkfed.org/api/rates/secured/sofr/last/1.json")
    if status != 200:
        raise RuntimeError("HTTP " + str(status))
    row = json.loads(body)["refRates"][0]
    when, value = row["effectiveDate"], str(row["percentRate"])
    show("rate", "SOFR", when, value, "percent", "business day")
    return when, value


def pink_sheet_table(rows):
    header = labels = None
    for index, row in enumerate(rows):
        text = [str(cell).strip() if cell else "" for cell in row]
        if "Crude oil, Brent" in text:
            header, labels = index, text
            break
    if labels is None:
        raise RuntimeError("Pink Sheet header not found")
    last = None
    for row in rows[header + 1:]:
        if row and re.fullmatch(r"\d{4}M\d{2}", str(row[0] or "")):
            last = row
    if last is None:
        raise RuntimeError("Pink Sheet has no month rows")
    updated = str(rows[3][0]) if len(rows) > 3 and rows[3] and rows[3][0] else ""
    return labels, last, updated


def fetch_worldbank():
    from openpyxl import load_workbook

    print("loading World Bank workbook...", flush=True)
    status, page = http_get("https://www.worldbank.org/en/research/commodity-markets")
    if status != 200:
        raise RuntimeError("page HTTP " + str(status))
    match = re.search(
        r"https://thedocs\.worldbank\.org/[^\"']+CMO-Historical-Data-Monthly\.xlsx",
        page.decode("utf-8", "replace"),
    )
    if not match:
        raise RuntimeError("workbook link missing")
    status, blob = http_get(match.group(0))
    if status != 200:
        raise RuntimeError("workbook HTTP " + str(status))
    rows = list(load_workbook(io.BytesIO(blob), read_only=True, data_only=True)["Monthly Prices"].iter_rows(values_only=True))
    labels, last, updated = pink_sheet_table(rows)
    found = {}
    for group, name, label, unit in WB:
        if label not in labels:
            raise RuntimeError(name + " column missing")
        value = str(last[labels.index(label)])
        found[name] = (str(last[0]), value)
        show(group, name, str(last[0]), value, unit, "monthly " + updated)
    return found


def daily_row(pulled_on, ecb, sofr, worldbank):
    sofr_as_of, sofr_value = sofr
    return [
        pulled_on,
        ecb["EURUSD"][1], ecb["EURUSD"][0],
        ecb["ECB_DEPOSIT"][1], ecb["ECB_DEPOSIT"][0],
        ecb["EURIBOR_3M"][1], ecb["EURIBOR_3M"][0],
        sofr_value, sofr_as_of,
        worldbank["BRENT"][1], worldbank["BRENT"][0],
        worldbank["GAS_EU"][1], worldbank["GAS_EU"][0],
        worldbank["MAIZE"][1], worldbank["MAIZE"][0],
    ]


def write_row(path, row):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = path.read_text(encoding="utf-8").splitlines() if path.exists() else []
    lines = [line for line in lines if line.strip()]
    if lines and lines[-1].split(",", 1)[0] == str(row[0]):
        lines.pop()
    buffer = io.StringIO()
    csv.writer(buffer, lineterminator="\n").writerow(row)
    lines.append(buffer.getvalue().rstrip("\n"))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def probe(path=None, pulled_on=None):
    path = Path(path) if path else output_path()
    pulled_on = pulled_on or date.today().isoformat()
    print(f"{'group':<8} {'series':<14} {'as_of':<12} {'value':<12} {'unit':<12} cadence", flush=True)
    try:
        row = daily_row(pulled_on, fetch_ecb(), fetch_sofr(), fetch_worldbank())
    except Exception as exc:
        print(f"skip     csv            {exc.__class__.__name__}: {exc}"[:160], flush=True)
        print("csv unchanged", flush=True)
        return 1
    write_row(path, row)
    print("wrote", path, flush=True)
    print(",".join(str(cell) for cell in row), flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(probe())
